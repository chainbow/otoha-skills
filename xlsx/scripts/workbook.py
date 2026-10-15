"""Builds Excel workbooks (.xlsx) from a JSON description, and reads them back.

    python scripts/workbook.py build --spec book.json --out 数据/预约.xlsx
    python scripts/workbook.py read  --file 数据/预约.xlsx

Paths are relative to the project folder, where the script runs. See the skill's SKILL.md for the sheets.
"""

import argparse
import datetime
import json
import os
import pathlib
import re
import sys

import openpyxl
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ACCENT = "5B57EF"
ACCENT_SOFT = "EEEDFE"
LINE = "E6E8EC"
FORMATS = {
    "text": "@",
    "number": "#,##0.00",
    "integer": "#,##0",
    "percent": "0.0%",
    "currency": "¥#,##0.00",
    "date": "yyyy-mm-dd",
}
CHARTS = ("column", "bar", "line", "pie")
MAX_READ_ROWS = 200
# The app keeps 64 KiB of what a script prints; what `read` prints stays under it.
MAX_OUTPUT = 60_000
MAX_READ_CELL = 500
ISO_DATE = re.compile(r"\d{4}-\d{2}-\d{2}([ T]\d{2}:\d{2}(:\d{2})?)?")


class WorkbookError(Exception):
    """The description cannot be made into a workbook; the message says what to change."""


def inside_project(path: str, must_exist: bool = False) -> pathlib.Path:
    project = pathlib.Path.cwd().resolve()
    candidate = (project / path).resolve()
    if candidate != project and project not in candidate.parents:
        raise WorkbookError(f"{path} is not inside the project folder.")
    if must_exist and not candidate.is_file():
        raise WorkbookError(f"{path} does not exist.")
    return candidate


def read_json(path: str):
    try:
        return json.loads(inside_project(path, must_exist=True).read_text(encoding="utf-8"))
    except UnicodeDecodeError as error:
        raise WorkbookError(f"{path} is not UTF-8 text; save it as UTF-8.") from error
    except json.JSONDecodeError as error:
        raise WorkbookError(f"{path} is not valid JSON: {error}") from error


def replaced(path: pathlib.Path, write):
    """Writes the file beside the old one and swaps it in, so a script stopped mid-save (the app ends
    scripts at their time, memory and write limits) leaves the old file whole. The name has no leading
    dot: the sandbox refuses those."""
    temporary = path.with_name(f"{path.stem}-saving-{os.getpid()}{path.suffix}")
    try:
        write(temporary)
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def described(value) -> str:
    if value is None:
        return "null"
    if isinstance(value, bool):
        return "true or false"
    if isinstance(value, (int, float)):
        return "a number"
    if isinstance(value, str):
        return "text"
    return "a list" if isinstance(value, list) else "an object"


def column_of(value, number):
    """A column: its header alone ("月份"), or an object with header, format, decimals and width."""
    if isinstance(value, (str, int, float)) and not isinstance(value, bool):
        return {"header": str(value)}
    if not isinstance(value, dict):
        raise WorkbookError(f'column {number} is a header, like "月份", or an object like {{"header": "月份", "format": "text"}}; it was {described(value)}.')
    return value


def number_format(column, index, header):
    kind = column.get("format", "text" if index == 1 else "number")
    if kind not in FORMATS:
        raise WorkbookError(f"format {kind!r} of {header!r} is not one of: {', '.join(FORMATS)}.")
    decimals = column.get("decimals")
    if decimals is None or kind != "number":
        return kind, FORMATS[kind]
    if isinstance(decimals, bool) or not isinstance(decimals, int) or not 0 <= decimals <= 6:
        raise WorkbookError(f"decimals of {header!r} is a whole number from 0 to 6; it was {json.dumps(decimals, ensure_ascii=False)}.")
    return kind, "#,##0" + ("." + "0" * decimals if decimals else "")


def put(cell, value, kind, where):
    """Writes one value. Text stays text even when it starts with "=": only {"formula": "…"} is a formula,
    so text copied from a page never runs in the user's spreadsheet."""
    if isinstance(value, dict):
        formula = value.get("formula")
        if set(value) != {"formula"} or not isinstance(formula, str) or not formula.strip():
            raise WorkbookError(f'{where}: a cell is text, a number, true or false, null, or {{"formula": "=SUM(B2:B4)"}}.')
        cell.value = formula if formula.startswith("=") else f"={formula}"
        return
    if isinstance(value, list):
        raise WorkbookError(f"{where}: a cell holds one value; it was a list.")
    if isinstance(value, str) and kind == "date" and ISO_DATE.fullmatch(value.strip()):
        cell.value = datetime.datetime.fromisoformat(value.strip())
        return
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def shown(value) -> str:
    return value["formula"] if isinstance(value, dict) and isinstance(value.get("formula"), str) else str(value)


def fill_sheet(sheet, spec, where):
    if not isinstance(spec.get("columns"), list) or not spec["columns"]:
        raise WorkbookError(f'{where}: a sheet needs columns, a list like ["月份", {{"header": "预约", "format": "integer"}}].')
    columns = [column_of(column, number) for number, column in enumerate(spec["columns"], start=1)]
    rows = spec.get("rows", [])
    if not isinstance(rows, list):
        raise WorkbookError(f"{where}: rows is a list of rows, each a list of cells; it was {described(rows)}.")
    headers = [str(column.get("header", "")) for column in columns]
    formats = []
    for index, column in enumerate(columns, start=1):
        try:
            formats.append(number_format(column, index, headers[index - 1]))
        except WorkbookError as error:
            raise WorkbookError(f"{where}: {error}") from error
    for number, row in enumerate(rows, start=1):
        if not isinstance(row, list) or len(row) != len(columns):
            raise WorkbookError(f"{where}, row {number}: it has {len(row) if isinstance(row, list) else 0} cells; the sheet has {len(columns)} columns.")

    thin = Side(style="thin", color=LINE)
    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=index)
        put(cell, header, "text", where)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=ACCENT)
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 22
    for number, row in enumerate(rows, start=2):
        for index, value in enumerate(row, start=1):
            cell = sheet.cell(row=number, column=index)
            put(cell, value, formats[index - 1][0], f"{where}, row {number - 1}, column {headers[index - 1]!r}")
            cell.number_format = formats[index - 1][1]
            cell.border = Border(bottom=thin)
    last = len(rows) + 1
    for index, column in enumerate(columns, start=1):
        width = column.get("width")
        if isinstance(width, bool) or not isinstance(width, (int, float)) or not 4 <= width <= 100:
            width = max(10, min(40, max((len(shown(value)) for value in [headers[index - 1]] + [row[index - 1] for row in rows]), default=10) * 1.6 + 2))
        sheet.column_dimensions[get_column_letter(index)].width = width

    totals = spec.get("totals")
    if totals:
        if not isinstance(totals, dict) or not isinstance(totals.get("sum", []), list):
            raise WorkbookError(f'{where}: totals is an object like {{"label": "合计", "sum": ["预约"]}}.')
        summed = totals.get("sum") or []
        unknown = [name for name in summed if name not in headers]
        if unknown:
            raise WorkbookError(f"{where}: totals name columns that do not exist: {', '.join(unknown)}.")
        put(sheet.cell(row=last + 1, column=1), str(totals.get("label", "合计")), "text", where)
        for name in summed:
            letter = get_column_letter(headers.index(name) + 1)
            sheet.cell(row=last + 1, column=headers.index(name) + 1).value = f"=SUM({letter}2:{letter}{last})"
        for index in range(1, len(headers) + 1):
            cell = sheet.cell(row=last + 1, column=index)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor=ACCENT_SOFT)
            cell.number_format = formats[index - 1][1]

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"

    chart_spec = spec.get("chart")
    if chart_spec:
        if not isinstance(chart_spec, dict):
            raise WorkbookError(f'{where}: chart is an object like {{"type": "column", "categories": "月份", "values": ["预约"]}}.')
        kind = chart_spec.get("type")
        if kind not in CHARTS:
            raise WorkbookError(f"{where}: chart type {kind!r} is not one of: {', '.join(CHARTS)}.")
        category = chart_spec.get("categories")
        values = chart_spec.get("values") or []
        if isinstance(values, str):
            values = [values]
        if category not in headers or not isinstance(values, list) or not values or any(name not in headers for name in values):
            raise WorkbookError(f"{where}: a chart names a column for its categories and one or more for its values.")
        chart = {"column": BarChart, "bar": BarChart, "line": LineChart, "pie": PieChart}[kind]()
        if kind == "bar":
            chart.type = "bar"
        chart.title = str(chart_spec["title"]) if chart_spec.get("title") else None
        chart.height, chart.width = 9, 16
        for name in values:
            column = headers.index(name) + 1
            chart.add_data(Reference(sheet, min_col=column, min_row=1, max_row=last), titles_from_data=True)
        column = headers.index(category) + 1
        chart.set_categories(Reference(sheet, min_col=column, min_row=2, max_row=last))
        sheet.add_chart(chart, f"{get_column_letter(len(headers) + 2)}2")


def build(arguments):
    spec = read_json(arguments.spec)
    sheets = spec.get("sheets") if isinstance(spec, dict) else None
    if not isinstance(sheets, list) or not sheets:
        raise WorkbookError('the description needs "sheets", a list of sheets.')
    out = inside_project(arguments.out)
    if out.suffix.lower() != ".xlsx":
        raise WorkbookError("--out must end in .xlsx.")
    book = openpyxl.Workbook()
    book.remove(book.active)
    for number, sheet_spec in enumerate(sheets, start=1):
        if not isinstance(sheet_spec, dict):
            raise WorkbookError(f'sheet {number} is an object like {{"name": "预约", "columns": [...], "rows": [...]}}; it was {described(sheet_spec)}.')
        # Excel refuses these characters in a sheet's name.
        name = re.sub(r"[\[\]:*?/\\]", "-", str(sheet_spec.get("name") or f"Sheet{number}")).strip("'")[:31] or f"Sheet{number}"
        fill_sheet(book.create_sheet(name), sheet_spec, f"sheet {name}")
    out.parent.mkdir(parents=True, exist_ok=True)
    replaced(out, book.save)
    print(json.dumps({"file": arguments.out, "sheets": [sheet.title for sheet in book.worksheets]}, ensure_ascii=False))


def readable(value):
    if isinstance(value, str):
        return value if len(value) <= MAX_READ_CELL else value[:MAX_READ_CELL] + "…"
    return value if isinstance(value, (int, float)) or value is None else str(value)


def read(arguments):
    path = inside_project(arguments.file, must_exist=True)
    try:
        book = openpyxl.load_workbook(path)
    except Exception as error:
        raise WorkbookError(f"{arguments.file} is not an Excel workbook that can be read ({error}).") from error
    sheets = []
    for sheet in book.worksheets:
        rows = [[readable(value) for value in row] for row in sheet.iter_rows(max_row=min(sheet.max_row, MAX_READ_ROWS), values_only=True)]
        entry = {"name": sheet.title, "rows": rows, "total_rows": sheet.max_row, "charts": len(sheet._charts)}
        if sheet.max_row > len(rows):
            entry["truncated"] = f"rows {len(rows) + 1} to {sheet.max_row} left out"
        sheets.append(entry)
    result = {"file": arguments.file, "sheets": sheets}
    # Too long for the app to pass on whole: the longest sheet gives up rows from its end until it fits.
    while len(json.dumps(result, ensure_ascii=False).encode()) > MAX_OUTPUT:
        longest = max(sheets, key=lambda entry: len(entry["rows"]))
        if len(longest["rows"]) <= 1:
            break
        longest["rows"] = longest["rows"][: len(longest["rows"]) * 3 // 4]
        longest["truncated"] = f"rows {len(longest['rows']) + 1} to {longest['total_rows']} left out: too long to show at once"
    print(json.dumps(result, ensure_ascii=False))


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    commands = parser.add_subparsers(dest="command", required=True)
    built = commands.add_parser("build")
    built.add_argument("--spec", required=True)
    built.add_argument("--out", required=True)
    reading = commands.add_parser("read")
    reading.add_argument("--file", required=True)
    arguments = parser.parse_args()
    try:
        {"build": build, "read": read}[arguments.command](arguments)
    except WorkbookError as error:
        print(f"workbook.py: {error}", file=sys.stderr)
        sys.exit(2)
    except (TypeError, ValueError, KeyError, AttributeError, OSError) as error:
        print(f"workbook.py: {type(error).__name__}: {error}", file=sys.stderr)
        sys.exit(2)


if __name__ == "__main__":
    main()
